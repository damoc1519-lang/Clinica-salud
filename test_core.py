from datetime import date,time,datetime,timedelta
import json
import unittest
from core import TZ,AppError,password_hash,password_ok,validate_patient,clinical_data,slots_for_day,BACKUP_TABLES
from exports import clinical_excel,clinical_pdf,billing_statement_pdf
from legacy import parse_schedule

class CoreTests(unittest.TestCase):
 def test_password_hash_is_salted_and_verifies(self):
    a=password_hash('ClaveSegura123'); b=password_hash('ClaveSegura123')
    assert a!=b and password_ok('ClaveSegura123',a) and not password_ok('otra',a)

 def test_patient_validation_normalizes_and_rejects_future_birth(self):
    p=validate_patient({'document':' 17 123 ','name':' Ana Ruiz ','birth_date':date(2000,1,1),'email':'a@b.ec'})
    assert p['document']=='17123' and p['name']=='Ana Ruiz'
    with self.assertRaises(AppError): validate_patient({'document':'1','name':'Ana','birth_date':date.today()+timedelta(days=1)})

 def test_clinical_data_keeps_habits_and_calculates_bmi(self):
    d=clinical_data({'habitos':'No fuma','motivo':'Dolor','diagnostico':'Control','peso':80,'talla':200,'fc':70},True)
    assert d['habitos']=='No fuma' and d['imc']==20

 def test_slots_remove_busy_blocked_and_past(self):
    day=date(2030,1,7) # lunes
    rule={'weekday':0,'start_time':time(8),'end_time':time(10)}
    busy={'start_at':datetime(2030,1,7,8,30,tzinfo=TZ),'end_at':datetime(2030,1,7,9,tzinfo=TZ)}
    block={'start_at':datetime(2030,1,7,9,tzinfo=TZ),'end_at':datetime(2030,1,7,9,30,tzinfo=TZ)}
    slots=slots_for_day(day,[rule],[busy],[block],30,datetime(2029,1,1,tzinfo=TZ))
    assert [s.strftime('%H:%M') for s in slots]==['08:00','09:30']

 def test_legacy_schedule_parser(self):
    days,start,end=parse_schedule('Lunes, Miércoles y Viernes','08:00 - 12:00')
    assert days==[0,2,4] and start==time(8) and end==time(12)

 def test_financial_tables_are_in_recoverable_backup(self):
    for table in ['services','agreement_tariffs','invoices','invoice_items','payments','cash_movements']:
       assert table in BACKUP_TABLES

 def test_internal_account_pdf(self):
    from decimal import Decimal
    inv={'number':'CTA-2026-000001','issue_date':date(2026,10,3),'due_date':date(2026,10,3),'patient':'ana ruiz','document':'1712345678','agreement':'ISSFA',
         'subtotal':Decimal('25'),'discount':Decimal('0'),'tax':Decimal('0'),'total':Decimal('25'),'paid':Decimal('5'),'patient_responsibility':Decimal('5'),'agreement_responsibility':Decimal('20'),'notes':'Prueba',
         'items':[{'description':'Consulta','quantity':Decimal('1'),'unit_price':Decimal('25'),'discount':Decimal('0'),'tax':Decimal('0'),'total':Decimal('25')}]}
    assert billing_statement_pdf(inv)[:4]==b'%PDF'

 def test_extended_clinical_record_and_exports(self):
    data=clinical_data({'motivo':'Control','diagnostico':'Hipertensión','peso':80,'talla':160,'temperatura':36.5,
                        'personal_conditions':['Hipertensión'],'lab_tests':['Biometría hemática'],'fum':date(2026,9,1),'prescription':'Losartán | 50 mg | 30 | 1 tableta | cada día | 30 días | 08:00'},True)
    assert data['imc']==31.25 and data['personal_conditions']==['Hipertensión']
    json.dumps(data)
    patient={'name':'ana ruiz','document':'1712345678','sex':'Femenino','birth_date':date(1980,1,1),'phone':'0999999999','address':'Quito','coverage':'Particular'}
    histories=[{'status':'Finalizada','occurred_at':datetime(2026,10,1,10,tzinfo=TZ),'doctor':'juan perez','specialty':'Medicina general','consultation_type':'C1','professional_id':'1700000000','registration':'MSP-001','data':data,'amendments':[]}]
    xlsx=clinical_excel(patient,histories)
    assert xlsx[:2]==b'PK'
    from openpyxl import load_workbook
    import io
    wb=load_workbook(io.BytesIO(xlsx),data_only=False)
    assert wb.sheetnames[:6]==['HC','INTER007','REF053','LAB010','IMA012','RECETA']
    assert wb['HC']['T3'].value=='MEDISUPORT / QMC'
    assert wb['HC']['AH3'].value=='1712345678'
    assert wb['HC']['B47'].value=='Hipertensión'
    assert wb['HC']['Y47'].value is None or wb['HC']['Y47'].value==''
    assert clinical_pdf(patient,histories)[:4]==b'%PDF'

if __name__=='__main__': unittest.main()
